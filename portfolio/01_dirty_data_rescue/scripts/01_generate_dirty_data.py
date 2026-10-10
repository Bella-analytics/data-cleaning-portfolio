"""
STEP 1 of 3 -- Build the messy, real-world-style source files.

This script intentionally creates BAD data. Its only purpose is to produce a
reproducible "before" state so that anyone can re-run this repository and see
the cleaning script work on the exact same input.

The three files mirror what clients actually send:
  - sales_orders_raw.csv      : a CRM / platform export (CSV, quoted fields)
  - customer_master_raw.csv   : a contact list with holes in it
  - inventory_jp_raw.xlsx     : a vendor spreadsheet with junk header/footer rows

Run:  python 01_generate_dirty_data.py
"""

from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

# ---------------------------------------------------------------- paths
BASE = Path(__file__).resolve().parent.parent
RAW = BASE / "data" / "raw"
RAW.mkdir(parents=True, exist_ok=True)

random.seed(20240613)  # fixed seed -> fully reproducible output

# ---------------------------------------------------------------- vocabulary
# Each "canonical" customer is written in several sloppy ways by real users.
CUSTOMERS = [
    ("Aurora Trading", ["Aurora Trading Ltd", "AURORA TRADING", "Aurora Trading  Ltd.",
                        "aurora trading ltd", " Aurora Trading", "Aurora Tradng Ltd"]),
    ("Meridian Supply", ["Meridian Supply Co", "MERIDIAN SUPPLY", "Meridian Supply Co.",
                         "Meridian  Supply", "meridian supply co"]),
    ("Northwind Retail", ["Northwind Retail", "NorthWind Retail Inc", "NORTHWIND RETAIL",
                          "Northwind Retail, Inc.", "northwind retail"]),
    ("Blue Harbor Foods", ["Blue Harbor Foods", "BLUE HARBOR FOODS", "Blue Harbour Foods",
                           "Blue Harbor Food", "blue harbor foods"]),
    ("Vertex Industrial", ["Vertex Industrial", "VERTEX INDUSTRIAL PTE", "Vertex Industrial Pte Ltd",
                           "Vertex Industral", "vertex industrial pte ltd"]),
    ("Sunfield Textiles", ["Sunfield Textiles", "SUNFIELD TEXTILES", "Sunfield Textile Co",
                           "Sunfield  Textiles", "sunfield textiles"]),
    ("Orion Logistics", ["Orion Logistics", "ORION LOGISTICS", "Orion Logistics Sdn Bhd",
                         "Orion Logistcs", "orion logistics"]),
    ("Pacific Rim Metals", ["Pacific Rim Metals", "PACIFIC RIM METALS", "Pacific Rim Metal",
                            "Pacific  Rim Metals", "pacific rim metals"]),
]

PRODUCTS = [
    ("Industrial Widget A1", ["industrial widget a1", "Industrial Widget  A1", "INDUSTRIAL WIDGET A1"]),
    ("Precision Gear B2", ["precision gear b2", "Precision  Gear B2", "PRECISION GEAR B2"]),
    ("Hydraulic Valve C3", ["hydraulic valve c3", "Hydraulic Valve C3 ", "HYDRAULIC VALVE C3"]),
    ("Sensor Module D4", ["sensor module d4", "Sensor  Module D4", "SENSOR MODULE D4"]),
    ("Steel Bracket E5", ["steel bracket e5", "Steel Bracket  E5", "STEEL BRACKET E5"]),
]

REGIONS = ["North", "South", "East", "West", "Central"]

REPS = ["J. Tan", "M. Rivera", "A. Kowalski", "S. Nakamura", "", "  ", "L. Chen"]

# amount is written the way a spreadsheet user would type it
AMOUNT_STYLES = [
    lambda v: f"{v:,.2f}",                 # 1,234.56
    lambda v: f"${v:,.2f}",                # $1,234.56
    lambda v: f"USD {v:,.2f}",             # USD 1,234.56
    lambda v: f"{v:.2f} USD",              # 1234.56 USD
    lambda v: f"'{v:,.2f}",                # leading apostrophe (Excel text artifact)
    lambda v: f"{v:.3f}",                  # stray extra decimal
    lambda v: f"({v:,.2f})",               # accounting-style negative
]

DATE_STYLES = [
    lambda d: d.isoformat(),                        # 2024-03-15
    lambda d: d.strftime("%d/%m/%Y"),               # 15/03/2024
    lambda d: d.strftime("%m/%d/%Y"),               # 03/15/2024  (ambiguous!)
    lambda d: d.strftime("%d-%b-%Y"),               # 15-Mar-2024
    lambda d: d.strftime("%Y/%m/%d"),               # 2024/03/15
    lambda d: d.strftime("%d.%m.%Y"),               # 15.03.2024
    lambda d: d.strftime("%b %d, %Y"),              # Mar 15, 2024
]

CURRENCIES = ["USD", "USD", "USD", "usd", "USD ", " USD", "$", "SGD", "MYR"]

JUNK = ["", "N/A", "n/a", "-", "null", "NULL", "None", "TBD", "#REF!", " "]


def maybe_junk(value: str, probability: float = 0.12) -> str:
    """Randomly replace a real value with the kind of placeholder users leave."""
    return random.choice(JUNK) if random.random() < probability else value


# ---------------------------------------------------------------- builder
def build_order_rows(n: int = 62) -> list[dict]:
    rows = []
    start = date(2024, 1, 6)
    for i in range(1, n + 1):
        canon_cust, variants = random.choice(CUSTOMERS)
        product, pvariants = random.choice(PRODUCTS)
        d = start + timedelta(days=random.randint(0, 300))
        qty = random.choice([1, 2, 3, 5, 8, 10, 12, 15, 20, 24, 30, 40, 50, 100, 0])
        unit = random.uniform(8, 950)
        amount = round(unit * max(qty, 1), 3)

        rows.append({
            "order_id": f"SO-{2024}-{i:04d}",
            "order_date": random.choice(DATE_STYLES)(d),
            "customer": random.choice(variants),
            "product": random.choice(pvariants),
            "qty": str(qty) if random.random() > 0.06 else maybe_junk(str(qty)),
            "amount": random.choice(AMOUNT_STYLES)(amount),
            "currency": random.choice(CURRENCIES),
            "status": random.choice(["completed", "Completed", "COMPLETE", "shipped",
                                     "Shipped", "pending", "Cancelled", "canceled", ""]),
            "region": maybe_junk(random.choice(REGIONS), 0.15),
            "rep": random.choice(REPS),
        })

    # --- inject duplicates: exact copies and near-duplicates -------------
    for src in random.sample(rows, 7):
        rows.append(dict(src))                                   # exact duplicate

    for src in random.sample(rows[:40], 4):
        near = dict(src)
        near["order_date"] = random.choice(DATE_STYLES)(
            date(2024, 1, 6) + timedelta(days=random.randint(0, 300))
        )                                                         # same id, new date
        near["amount"] = random.choice(AMOUNT_STYLES)(random.uniform(50, 900))
        rows.append(near)

    # --- inject whitespace / invisible-character damage -------------------
    for row in rows:
        if random.random() < 0.30:
            row["customer"] = "  " + row["customer"] + " "
        if random.random() < 0.18:
            row["product"] = row["product"] + "\u00a0"            # non-breaking space
        if random.random() < 0.10:
            row["region"] = row["region"].upper() if row["region"] else row["region"]

    # --- blank rows, the classic extra empty line -------------------------
    for _ in range(4):
        rows.insert(random.randint(0, len(rows)), {k: "" for k in rows[0]})

    random.shuffle(rows)
    return rows


def build_customer_rows() -> list[dict]:
    rows = []
    for canon, variants in CUSTOMERS:
        for v in random.sample(variants, k=random.randint(1, 3)):
            phone = random.choice([
                f"+65 {random.randint(6000,9999)} {random.randint(1000,9999)}",
                f"({random.randint(200,999)}) {random.randint(200,999)}-{random.randint(1000,9999)}",
                f"{random.randint(200,999)}.{random.randint(200,999)}.{random.randint(1000,9999)}",
                f"+1-{random.randint(200,999)}-{random.randint(200,999)}-{random.randint(1000,9999)}",
                "N/A", "", "  ",
            ])
            rows.append({
                "customer_name": v,
                "contact_email": maybe_junk(
                    f"ops@{canon.split()[0].lower()}.com", 0.20
                ),
                "phone": phone,
                "country": maybe_junk(random.choice(["SG", "US", "MY", "sg", "United States", ""]), 0.15),
                "credit_limit": random.choice(
                    [f"{random.randint(1,50)*1000:,}", "N/A", "", f"USD {random.randint(1,50)*1000:,}"]
                ),
            })
    random.shuffle(rows)
    return rows


def build_inventory_rows() -> list[dict]:
    names = ["精密ギア B2", "油圧バルブ C3", "センサーモジュール D4", "産業用ウィジェット A1"]
    rows = []
    for i, nm in enumerate(names * 3, start=1):
        rows.append({
            "sku": f"INV-{i:03d}",
            "name_jp": nm,
            "stock": random.choice([str(random.randint(0, 500)), "在庫切れ", "", "unknown"]),
            "unit_price_jpy": random.choice([
                f"¥{random.randint(1000,90000):,}",
                f"{random.randint(1000,90000):,}円",
                f"{random.randint(1000,90000)}",
                "N/A",
            ]),
            "note": random.choice(["", "  ", "要確認", "廃番", None]),
        })
    return rows


# ---------------------------------------------------------------- writers
def write_csv(path: Path, rows: list[dict]) -> None:
    # Header text carries stray spaces, exactly as exported by many tools.
    fieldnames = [f"  {k} " if i % 2 == 0 else f"{k}  " for i, k in enumerate(rows[0].keys())]
    pretty = [k.strip() for k in fieldnames]
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, quoting=csv.QUOTE_ALL)
        writer.writerow(fieldnames)
        for row in rows:
            writer.writerow([row.get(k, "") for k in pretty])
    print(f"  wrote {path.name}: {len(rows)} rows, header={fieldnames}")


def write_inventory_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "在庫"

    # Junk rows above the real table -- the single most common Excel problem.
    ws["A1"] = "VENDOR STOCK REPORT"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "generated automatically - do not edit"
    ws["A3"] = ""            # blank spacer row
    ws["A4"] = "Currency: JPY"

    header_row = 6
    headers = ["sku", "name_jp", "stock", "unit_price_jpy", "note"]
    for col, name in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=name)

    for r, row in enumerate(rows, start=header_row + 1):
        for c, key in enumerate(headers, start=1):
            ws.cell(row=r, column=c, value=row[key])

    # Footer junk rows after the table.
    last = header_row + len(rows)
    ws.cell(row=last + 2, column=1, value="--- END OF REPORT ---")
    ws.cell(row=last + 3, column=1, value="prepared by: warehouse team")
    ws.cell(row=last + 4, column=1, value="total SKUs:")
    ws.cell(row=last + 4, column=3, value=len(rows))

    wb.save(path)
    print(f"  wrote {path.name}: table starts at row {header_row}, "
          f"{len(rows)} data rows, then {3} junk footer rows")


def main() -> None:
    print("Building dirty source files in", RAW)
    orders = build_order_rows()
    write_csv(RAW / "sales_orders_raw.csv", orders)
    write_csv(RAW / "customer_master_raw.csv", build_customer_rows())
    write_inventory_xlsx(RAW / "inventory_jp_raw.xlsx", build_inventory_rows())
    print("Done. These files are SUPPOSED to look wrong.")


if __name__ == "__main__":
    main()
